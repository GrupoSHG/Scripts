"""
01_restaurar_y_extraer.py
=========================
1. Restaura el .bak más reciente en SQL Server
2. Lee y limpia múltiples archivos .sql desde una carpeta
3. Exporta el resultado a un único Excel con múltiples pestañas
"""

import os
import shutil
import time
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Dict

import pyodbc
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
#  CONFIGURACIÓN
# ─────────────────────────────────────────────
SQL_SERVER     = os.environ.get("SQL_SERVER", r"localhost\SQLEXPRESS")
# En tu notebook, sigue usando autenticación de Windows (Trusted_Connection)
# igual que siempre — no hace falta tocar nada localmente. En GitHub Actions,
# la instancia de SQL Server se instala sin nombre y con autenticación SQL
# (usuario 'sa' + password), así que ahí se setean estas 2 variables de
# entorno para activar ese modo.
SQL_USE_SQL_AUTH = os.environ.get("SQL_USE_SQL_AUTH", "false").lower() == "true"
SQL_SA_PASSWORD  = os.environ.get("SQL_SA_PASSWORD", "")

def _sql_conn_str(database):
    if SQL_USE_SQL_AUTH:
        return (f"DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={SQL_SERVER};"
                f"DATABASE={database};UID=sa;PWD={SQL_SA_PASSWORD};")
    return (f"DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={SQL_SERVER};"
            f"DATABASE={database};Trusted_Connection=yes;")

NOMBRE_BD      = "T779354202C"
CARPETA_BAK    = r"C:\Backups\Manager"
CARPETA_SALIDA = r"C:\Reportes\VentasFull"
CARPETA_SQL    = Path(__file__).parent / "ConsultasSQL" # <--- CARPETA DONDE PONDRÁS TUS ARCHIVOS .SQL
DATA_DIR       = r"C:\Program Files\Microsoft SQL Server\MSSQL17.SQLEXPRESS\MSSQL\DATA"
AÑO_ACTUAL     = datetime.now().year
# ─────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)s  %(message)s",
    handlers=[
        logging.FileHandler(r"C:\Users\atorr\Documents\pipeline_ventas.log", encoding="utf-8"),
        logging.StreamHandler()
    ]
)
log = logging.getLogger(__name__)

# Caracteres de control ilegales para el formato XML interno de .xlsx
# (excluye tab \x09, salto de línea \x0A y retorno de carro \x0D, que sí
# están permitidos). Suelen venir de columnas TEXT viejas de SQL Server.
_RE_CARACTERES_ILEGALES = re.compile('[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]')

def limpiar_caracteres_ilegales(valor):
    """Quita caracteres de control que openpyxl/Excel rechazan."""
    if not isinstance(valor, str):
        return valor
    return _RE_CARACTERES_ILEGALES.sub('', valor)


def conectar_master() -> pyodbc.Connection:
    conn_str = _sql_conn_str("master")
    return pyodbc.connect(conn_str, autocommit=True)

def obtener_bak_reciente() -> Path:
    ultimo_txt = Path("ultimo_bak.txt")
    if ultimo_txt.exists():
        ruta = Path(ultimo_txt.read_text().strip())
        if ruta.exists():
            log.info(f"Usando .bak de ultimo_bak.txt: {ruta.name}")
            return ruta

    archivos = list(Path(CARPETA_BAK).glob("*.bak"))
    if not archivos:
        raise FileNotFoundError(f"No se encontraron archivos .bak en {CARPETA_BAK}")
    return max(archivos, key=lambda f: f.stat().st_mtime)

def detectar_data_dir() -> str:
    if Path(DATA_DIR).exists(): return DATA_DIR
    bases = [r"C:\Program Files\Microsoft SQL Server", r"C:\Program Files (x86)\Microsoft SQL Server"]
    for base in bases:
        base_path = Path(base)
        if not base_path.exists(): continue
        for carpeta in sorted(base_path.iterdir(), reverse=True):
            data = carpeta / "MSSQL" / "DATA"
            if data.exists():
                return str(data)
    raise FileNotFoundError("No se encontró la carpeta DATA de SQL Server.")

def restaurar_bd(bak_path: Path):
    log.info(f"Restaurando {bak_path.name} como '{NOMBRE_BD}'...")

    data_dir = detectar_data_dir()
    bak_en_data = Path(data_dir) / bak_path.name
    if not bak_en_data.exists():
        log.info(f"Copiando .bak a {bak_en_data}...")
        shutil.copy2(str(bak_path), str(bak_en_data))
    
    conn = conectar_master()
    cursor = conn.cursor()

    cursor.execute(f"IF EXISTS (SELECT name FROM sys.databases WHERE name = N'{NOMBRE_BD}') ALTER DATABASE [{NOMBRE_BD}] SET SINGLE_USER WITH ROLLBACK IMMEDIATE")

    cursor.execute(f"RESTORE FILELISTONLY FROM DISK = N'{bak_en_data}'")
    files = cursor.fetchall()
    logical_data = next((r[0] for r in files if r[2].upper() == 'D'), None)
    logical_log  = next((r[0] for r in files if r[2].upper() == 'L'), None)

    data_file = os.path.join(data_dir, f"{NOMBRE_BD}.mdf")
    log_file  = os.path.join(data_dir, f"{NOMBRE_BD}_log.ldf")

    log.info("Ejecutando RESTORE DATABASE...")
    cursor.execute(f"""
        RESTORE DATABASE [{NOMBRE_BD}]
        FROM DISK = N'{bak_en_data}'
        WITH MOVE N'{logical_data}' TO N'{data_file}', MOVE N'{logical_log}'  TO N'{log_file}',
        REPLACE, RECOVERY, STATS = 10
    """)

    for _ in range(15):
        time.sleep(3)
        try:
            check = conectar_master()
            row = check.cursor().execute(f"SELECT state_desc FROM sys.databases WHERE name = N'{NOMBRE_BD}'").fetchone()
            check.close()
            if row and row[0] == 'ONLINE': break
        except Exception: pass

    for _ in range(5):
        try:
            time.sleep(2)
            c2 = conectar_master()
            c2.cursor().execute(f"ALTER DATABASE [{NOMBRE_BD}] SET MULTI_USER")
            c2.close()
            break
        except Exception: pass
    
    conn.close()
    bak_en_data.unlink(missing_ok=True)
    log.info(f"✅ BD '{NOMBRE_BD}' restaurada correctamente.")

def extraer_reportes() -> Dict[str, pd.DataFrame]:
    """ Lee todos los .sql de la carpeta, los limpia y los ejecuta """
    log.info(f"Conectando a '{NOMBRE_BD}' para extraer reportes...")
    conn_str = _sql_conn_str(NOMBRE_BD)
    conn = pyodbc.connect(conn_str)
    
    reportes_df = {}
    carpeta = Path(CARPETA_SQL)
    
    if not carpeta.exists():
        raise FileNotFoundError(f"No existe la carpeta de consultas: {CARPETA_SQL}")

    for archivo_sql in carpeta.glob("*.sql"):
        nombre_reporte = archivo_sql.stem # Nombre sin extensión
        log.info(f"Ejecutando: {nombre_reporte}...")
        
        with open(archivo_sql, 'r', encoding='utf-8') as f:
            sql_text = f.read()

        # 1. Reemplazar variables de Python
        sql_text = sql_text.replace("{AÑO_ACTUAL}", str(AÑO_ACTUAL))
        
        # 2. Limpieza estricta para PyODBC
        # Eliminar comandos 'GO' (PyODBC no los soporta)
        sql_text = re.sub(r'(?i)^\s*GO\s*$', '', sql_text, flags=re.MULTILINE)
        # Eliminar 'USE T...' (Obligamos a que use la BD conectada)
        sql_text = re.sub(r'(?i)^\s*USE\s+\[?.*?\]?\s*;?', '', sql_text, flags=re.MULTILINE)
        
        # 3. SET NOCOUNT ON es obligatorio para consultas con DECLARE
        sql_listo = f"SET NOCOUNT ON;\n{sql_text}"
        
        try:
            df = pd.read_sql(sql_listo, conn)
            reportes_df[nombre_reporte] = df
            log.info(f"  └─ {len(df):,} filas extraídas.")
        except Exception as e:
            log.error(f"  └─ Error ejecutando {nombre_reporte}: {e}")

    conn.close()
    return reportes_df

def exportar_excel(diccionario_dfs: Dict[str, pd.DataFrame]) -> Path:
    Path(CARPETA_SALIDA).mkdir(parents=True, exist_ok=True)
    fecha_hoy = datetime.now().strftime("%Y-%m-%d")
    ruta_excel = Path(CARPETA_SALIDA) / f"Reportes_Completos_{fecha_hoy}.xlsx"

    log.info(f"Exportando a Excel multipestaña: {ruta_excel}")

    with pd.ExcelWriter(ruta_excel, engine="openpyxl") as writer:
        for nombre_reporte, df in diccionario_dfs.items():
            if df.empty:
                log.warning(f"  └─ {nombre_reporte}: DataFrame vacío, omitido")
                continue
            
            # 🔹 Renombrar columnas duplicadas (causa del error original)
            if df.columns.duplicated().any():
                cols = pd.Series(df.columns)
                for dup in cols[cols.duplicated()].unique():
                    indices = cols[cols == dup].index.tolist()
                    for i, idx in enumerate(indices):
                        if i > 0:
                            cols[idx] = f"{dup}_{i}"
                df.columns = cols.tolist()
                log.warning(f"  └─ {nombre_reporte}: columnas duplicadas renombradas")
                
            # Excel limita los nombres de pestaña a 31 caracteres
            nombre_pestana = nombre_reporte[:31]

            # 🔹 Limpiar caracteres de control ilegales para el XML interno
            # de .xlsx (residuo típico de columnas TEXT viejas de SQL Server,
            # invisibles a simple vista pero que openpyxl rechaza)
            for col in df.select_dtypes(include='object').columns:
                df[col] = df[col].apply(limpiar_caracteres_ilegales)

            df.to_excel(writer, sheet_name=nombre_pestana, index=False)
            ws = writer.sheets[nombre_pestana]

            # Estilos
            header_fill = PatternFill("solid", start_color="2F5496")
            header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            thin        = Side(style="thin", color="CCCCCC")
            border      = Border(left=thin, right=thin, bottom=thin)

            # Formatear encabezados
            for col_idx in range(1, len(df.columns) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            ws.row_dimensions[1].height = 25

            # 🔹 Formato dinámico de celdas — iteramos por POSICIÓN, no por nombre
            for col_idx in range(len(df.columns)):
                col_name = df.columns[col_idx]
                col_name_upper = str(col_name).upper()
                es_pesos = any(kw in col_name_upper for kw in ["TOTAL", "PESO", "PRECIO", "CTO", "VALOR", "PEND", "PAGADO", "FACTURADO"])
                es_cantidad = any(kw in col_name_upper for kw in ["CANTIDAD", "Q ", "WIP", "PRODUCCION"])
                es_pct = any(kw in col_name_upper for kw in ["PCT", "PORCENTAJE", "DESCTO"])

                # 🔹 Acceder por iloc (índice posicional), nunca por nombre
                try:
                    serie = df.iloc[:, col_idx]
                    if len(serie) > 0:
                        max_len_datos = serie.astype(str).str.len().max()
                    else:
                        max_len_datos = 0
                except Exception as e:
                    log.warning(f"     Col {col_idx} ({col_name}): no se pudo medir ancho ({e})")
                    max_len_datos = 10
                
                max_len = max(len(str(col_name)), max_len_datos)
                col_letter = get_column_letter(col_idx + 1)
                ws.column_dimensions[col_letter].width = min(max_len + 2, 45)

                # Formato de filas
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx + 1, max_col=col_idx + 1):
                    for cell in row:
                        cell.font = Font(name="Arial", size=9)
                        cell.border = border
                        
                        if pd.notna(cell.value) and isinstance(cell.value, (int, float)):
                            if es_pct:
                                cell.number_format = '0.00%'
                            elif es_cantidad:
                                cell.number_format = '#,##0.00'
                            elif es_pesos:
                                cell.number_format = '#,##0'

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

    log.info(f"✅ Excel generado: {ruta_excel} ({ruta_excel.stat().st_size/1e6:.1f} MB)")
    return ruta_excel



def main():
    log.info("=" * 60)
    log.info("  PASO 1 — Restaurar BD y extraer Múltiples Reportes")
    log.info("=" * 60)

    bak = obtener_bak_reciente()
    log.info(f"Usando backup: {bak.name}")

    restaurar_bd(bak)
    
    # Extrae todos los reportes de la carpeta SQL
    diccionario_reportes = extraer_reportes()
    
    # Crea el Excel multipestaña
    ruta_excel = exportar_excel(diccionario_reportes)

    with open("ultimo_excel.txt", "w") as f:
        f.write(str(ruta_excel))

    log.info(f"✅ Proceso Completado → {ruta_excel}")
    return ruta_excel

if __name__ == "__main__":
    main()